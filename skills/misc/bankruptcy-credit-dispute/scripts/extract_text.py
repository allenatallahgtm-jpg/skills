#!/usr/bin/env python3
"""Extract text from credit reports and bankruptcy papers (PDF, DOCX, TXT/HTML).

Usage:
    python extract_text.py <file-or-dir> [more ...] --out <dir>

Writes one <name>.txt per input file into --out and prints a summary.
PDF pages with no extractable text (scanned images) are listed so the agent
can read those pages visually instead.
"""
import argparse
import html
import re
import sys
from pathlib import Path

SUPPORTED = {".pdf", ".docx", ".txt", ".html", ".htm"}


def extract_pdf(path: Path):
    try:
        from pypdf import PdfReader
    except ImportError:
        sys.exit("pypdf missing. Run: python3 -m venv /tmp/venv && /tmp/venv/bin/pip install pypdf python-docx")
    reader = PdfReader(str(path))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            return "", ["ENCRYPTED - ask user for password or an unlocked copy"]
    parts, empty = [], []
    for i, page in enumerate(reader.pages, 1):
        text = (page.extract_text() or "").strip()
        if len(text) < 20:
            empty.append(i)
        parts.append(f"\n===== PAGE {i} =====\n{text}")
    notes = [f"pages with no text layer (scanned? read visually): {empty}"] if empty else []
    return "".join(parts), notes


def extract_docx(path: Path):
    try:
        import docx
    except ImportError:
        sys.exit("python-docx missing. Run: /tmp/venv/bin/pip install python-docx")
    d = docx.Document(str(path))
    lines = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(lines), []


def extract_html(path: Path):
    raw = path.read_text(errors="ignore")
    raw = re.sub(r"(?is)<(script|style).*?</\1>", " ", raw)
    raw = re.sub(r"(?i)<br\s*/?>|</(p|div|tr|li|h\d)>", "\n", raw)
    raw = re.sub(r"(?i)</t[dh]>", " | ", raw)
    text = html.unescape(re.sub(r"<[^>]+>", " ", raw))
    return re.sub(r"[ \t]+", " ", text), []


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("inputs", nargs="+")
    ap.add_argument("--out", required=True)
    args = ap.parse_args()
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)

    files = []
    for item in args.inputs:
        p = Path(item)
        files += sorted(f for f in p.rglob("*") if f.suffix.lower() in SUPPORTED) if p.is_dir() else [p]

    for f in files:
        ext = f.suffix.lower()
        if ext == ".pdf":
            text, notes = extract_pdf(f)
        elif ext == ".docx":
            text, notes = extract_docx(f)
        elif ext in (".html", ".htm"):
            text, notes = extract_html(f)
        elif ext == ".txt":
            text, notes = f.read_text(errors="ignore"), []
        else:
            print(f"SKIP {f} (unsupported; if it is an image, read it visually)")
            continue
        dest = out / (f.stem + ".txt")
        dest.write_text(text)
        print(f"OK   {f} -> {dest} ({len(text)} chars){' | ' + '; '.join(notes) if notes else ''}")


if __name__ == "__main__":
    main()
