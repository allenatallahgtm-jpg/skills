#!/usr/bin/env python3
"""Analyze a bankruptcy case file against credit reports and build the dispute packet.

Usage:
    python generate_packet.py <case.json> --out <dir> [--today YYYY-MM-DD]

Input format: see ../references/case.example.json
Outputs (in --out), one issue per letter, each letter as .md and .docx:
    analysis.md                              findings, grounds, items needing attorney review
    letters/<bureau>-obsolete                public record or accounts past the reporting period
    letters/<bureau>-public-record           inaccurate, incomplete, duplicate, or mixed-file public record
    letters/<bureau>-discharged-accounts     discharged accounts still reported as owing
    letters/<bureau>-identity-theft-block    only when options.identity_theft_confirmed + ftc_report_number
    letters/<recipient>-file-disclosure      file disclosure request (missing reports or options.request_disclosure)
    letters/lexisnexis-vendor-dispute        public-records vendor dispute
    letters/furnisher-*                      direct disputes to furnishers (15 U.S.C. 1681s-2(a)(8), 12 C.F.R. 1022.43)
    client-authorization                     attorney signer only
    checklist.md                             mailing checklist, addresses, response deadlines
Only grounds supported by the data are asserted. Anything uncertain is listed
under "Needs attorney review" and is NOT put in a letter.
"""
import argparse
import json
import re
import sys
from datetime import date, timedelta
from pathlib import Path

# Addresses: verify before mailing (see checklist).
BUREAUS = {
    "equifax": ("Equifax Information Services LLC", "P.O. Box 740256", "Atlanta, GA 30374-0256"),
    "experian": ("Experian", "P.O. Box 4500", "Allen, TX 75013"),
    "transunion": ("TransUnion Consumer Solutions", "P.O. Box 2000", "Chester, PA 19016-2000"),
    "innovis": ("Innovis Consumer Assistance", "P.O. Box 26", "Pittsburgh, PA 15230-0026"),
}
RECIPIENTS = {**BUREAUS,
              "lexisnexis": ("LexisNexis Risk Solutions Consumer Center", "P.O. Box 105108", "Atlanta, GA 30348-5108")}
DISPLAY = {"equifax": "Equifax", "experian": "Experian", "transunion": "TransUnion", "innovis": "Innovis",
           "lexisnexis": "LexisNexis Risk Solutions"}
NATIONWIDE = ("equifax", "experian", "transunion")
BK_STATUS_WORDS = ("bankrupt", "discharg", "chapter", "ch 7", "ch 13", "included")
DELETE_KINDS = ("obsolete", "ch13_policy", "not_client", "unverifiable")
MAIL_DAYS = 5


# ---------------------------------------------------------------- helpers
def d(s):
    return date.fromisoformat(s) if s else None


def add_years(dt, n):
    try:
        return dt.replace(year=dt.year + n)
    except ValueError:  # Feb 29
        return dt.replace(year=dt.year + n, day=28)


def add_business_days(dt, n):
    """Weekends skipped; federal holidays are not (checklist says so)."""
    while n > 0:
        dt += timedelta(days=1)
        if dt.weekday() < 5:
            n -= 1
    return dt


def norm(s):
    return re.sub(r"[^a-z0-9]", "", str(s or "").lower())


def norm_case(s):
    # "2:15-bk-12345-ABC" vs "15-12345" -> compare year+number core
    digits = re.findall(r"\d+", str(s or ""))
    digits = [x for x in digits if len(x) >= 2]
    if len(digits) >= 2:
        return digits[-2][-2:] + digits[-1].lstrip("0")
    return norm(s)


def norm_status(s):
    s = norm(s)
    for key in ("dismiss", "withdraw", "discharg", "open", "filed", "closed"):
        if key in s:
            return {"withdraw": "withdrawn", "dismiss": "dismissed", "discharg": "discharged"}.get(key, key)
    return s


def money(v):
    try:
        return float(str(v).replace("$", "").replace(",", "") or 0)
    except ValueError:
        return 0.0


def fmt(dt):
    return dt.strftime("%B %-d, %Y") if dt else "[date]"


def slug(s):
    return re.sub(r"[^a-z0-9]+", "-", str(s).lower()).strip("-")


def records(pr):
    """public_record may be null, one object, or a list of objects."""
    if not pr:
        return []
    return [x for x in (pr if isinstance(pr, list) else [pr]) if x]


def signer(case):
    return str(case.get("signer") or "attorney").lower()


def v(case, attorney_text, consumer_text):
    return attorney_text if signer(case) == "attorney" else consumer_text


def name_of(recipient):
    return DISPLAY.get(recipient, recipient.title())


def reinvestigation_days(case, bureau, extended=False):
    """45 days after a free annual disclosure (1681j(a)(3)) or when extended (1681i(a)(1)(B)); else 30."""
    src = (case.get("credit_reports", {}).get(bureau) or {}).get("report_source", "other")
    return 45 if src == "annual_free" or extended else 30


def deadline_basis(case, bureau, extended=False):
    src = (case.get("credit_reports", {}).get(bureau) or {}).get("report_source", "other")
    if src == "annual_free":
        return "15 U.S.C. § 1681j(a)(3)"
    if extended:
        return "15 U.S.C. § 1681i(a)(1)(A) and (B)"  # 30 days plus an extension of not more than 15
    return "15 U.S.C. § 1681i(a)(1)(A)"


def is_lexis(s):
    return bool(re.search(r"lexis", str(s or ""), re.I))


# ---------------------------------------------------------------- analysis
def analyze(case, today):
    court = case["bankruptcy_court_record"]
    client = case["client"]
    opts = case.get("options", {})
    chapter = str(court.get("chapter", "")).strip()
    filed = d(court.get("filing_date"))
    disposition = norm_status(court.get("disposition"))
    disp_date = d(court.get("disposition_date"))

    excluded = {(norm(x.get("account_last4")), norm(x.get("creditor"))): x for x in court.get("excluded_from_discharge", [])}
    scheduled = court.get("scheduled_creditors", [])
    idt_confirmed = bool(opts.get("identity_theft_confirmed"))
    idt_ok = idt_confirmed and bool(opts.get("ftc_report_number"))

    results, review = {}, []
    if not filed:
        review.append("Court filing date missing: cannot evaluate obsolescence. Get the petition / docket from PACER.")
    if idt_confirmed and not idt_ok:
        review.append("options.identity_theft_confirmed is set but options.ftc_report_number is missing: no identity theft "
                      "block letter written. File the FTC Identity Theft Report at IdentityTheft.gov first.")

    for bureau, rep in case.get("credit_reports", {}).items():
        bureau = bureau.lower()
        if bureau not in BUREAUS:
            review.append(f"Unknown bureau '{bureau}' skipped.")
            continue
        rep = rep or {}
        B = name_of(bureau)
        recs = records(rep.get("public_record"))
        r = {"obsolete": [], "obsolete_tradelines": [], "public_record": [], "reserve": [], "tradelines": [],
             "identity_theft": False, "vendor_source": None, "records": recs, "notes": []}

        if recs:
            # Obsolete (statutory) or Chapter 13 seven-year bureau policy, computed from the court filing date.
            if filed and today >= add_years(filed, 10):
                r["obsolete"].append(
                    ("obsolete", f"**Obsolete information.** The order for relief was entered on {fmt(filed)}, more than ten years ago. "
                     "Under 15 U.S.C. § 1681c(a)(1), this item may no longer be reported."))
            elif chapter == "13" and filed and today >= add_years(filed, 7) and opts.get("include_chapter13_policy", True):
                # Bureau policy, not statute: Equifax, Experian, and TransUnion publish it; Innovis does not.
                policy = ("Consistent with your published practice of removing Chapter 13 bankruptcies seven years from the filing date, "
                          "please delete it." if bureau in NATIONWIDE else
                          "Equifax, Experian, and TransUnion remove Chapter 13 bankruptcies seven years from the filing date as a "
                          "matter of policy. Please apply the same practice and delete it.")
                r["obsolete"].append(
                    ("ch13_policy", f"**Chapter 13 reporting period expired.** This Chapter 13 case was filed on {fmt(filed)}, more than seven years ago. "
                     + policy))

            grounds = []
            multi = len(recs) > 1
            id_issue_all = []
            for i, pr in enumerate(recs, 1):
                tag = f"Record {i} of {len(recs)}: " if multi else ""
                mism = []
                chapter_wrong = bool(pr.get("chapter") and chapter and norm(pr["chapter"]) != norm(chapter))
                if chapter_wrong:
                    mism.append(f"chapter (reported: {pr['chapter']}; court record: {chapter})")
                if pr.get("case_number") and court.get("case_number") and norm_case(pr["case_number"]) != norm_case(court["case_number"]):
                    mism.append(f"case number (reported: {pr['case_number']}; court record: {court['case_number']})")
                if pr.get("filing_date") and filed and d(pr["filing_date"]) != filed:
                    mism.append(f"filing date (reported: {fmt(d(pr['filing_date']))}; court record: {fmt(filed)})")
                if pr.get("status_date") and disp_date and d(pr["status_date"]) != disp_date:
                    mism.append(f"{disposition or 'status'} date (reported: {fmt(d(pr['status_date']))}; court record: {fmt(disp_date)})")
                rep_status = norm_status(pr.get("status"))
                if rep_status and disposition and rep_status != disposition and not (rep_status == "closed" and disposition == "discharged"):
                    if disposition in ("dismissed", "withdrawn"):
                        text = (f"**{tag}{disposition.title()} case reported incorrectly.** The court record shows the case was "
                                f"{disposition} on {fmt(disp_date)}, but it is reported as \"{pr.get('status')}\".")
                        if disposition == "withdrawn":
                            text += (" If a bankruptcy case is withdrawn by the consumer before a final judgment, 15 U.S.C. § 1681c(d)(1) "
                                     "requires the report to state that the case was withdrawn, upon receipt of documentation certifying "
                                     "the withdrawal. The enclosed court record (Exhibit C) is that documentation.")
                        grounds.append(("disposition", text))
                    else:
                        mism.append(f"status (reported: {pr.get('status')}; court record: {disposition} on {fmt(disp_date)})")
                if mism:
                    text = (f"**{tag}Inaccurate case details.** The record misstates the " + "; ".join(mism) +
                            ". The certified court records are enclosed as Exhibit B.")
                    if chapter_wrong:
                        text += (" 15 U.S.C. § 1681c(d)(1) requires a report containing a bankruptcy case to identify the chapter "
                                 "under which the case arises, if provided by the source of the information; the chapter shown is wrong.")
                    grounds.append(("inaccurate", text))

                # Identity mismatch -> review unless attorney confirmed
                id_issue = []
                if pr.get("ssn_last4") and norm(pr["ssn_last4"]) != norm(client.get("ssn_last4")):
                    id_issue.append(f"SSN last 4 on record ({pr['ssn_last4']}) does not match client")
                names = [norm(client.get("name"))] + [norm(a) for a in client.get("aliases", [])]
                if pr.get("name") and norm(pr["name"]) not in names:
                    id_issue.append(f"name on record ('{pr['name']}') does not match client or aliases")
                id_issue_all += [tag + x for x in id_issue]
                if is_lexis(pr.get("source")):
                    r["vendor_source"] = pr.get("source")

            if court.get("debtor_name") and norm(court["debtor_name"]) not in \
                    [norm(client.get("name"))] + [norm(a) for a in client.get("aliases", [])]:
                id_issue_all.append(f"court debtor name ('{court['debtor_name']}') does not match client: confirm this is the client's case")
            if id_issue_all:
                if opts.get("not_client_record_confirmed"):
                    grounds.append(
                        ("not_client", v(case, "**Not my client's record.** My client did not file the bankruptcy reported. ",
                                         "**Not my record.** I did not file the bankruptcy reported. ") +
                         "This appears to be a mixed-file or identity error: " + "; ".join(id_issue_all) + "."))
                else:
                    review.append(f"{B}: possible mixed file: " + "; ".join(id_issue_all) +
                                  ". Set options.not_client_record_confirmed=true only if the client confirms they never filed.")

            # Duplicate listing of the same case
            core = norm_case(court.get("case_number")) if court.get("case_number") else None
            same = [pr for pr in recs if (core and pr.get("case_number") and norm_case(pr["case_number"]) == core)
                    or (not pr.get("case_number") and filed and d(pr.get("filing_date")) == filed)]
            if len(same) > 1:
                grounds.append(
                    ("duplicate", f"**Duplicate listing.** The same bankruptcy case (Case No. {court.get('case_number')}) is reported "
                     f"{len(same)} times, which makes it appear that {v(case, 'my client', 'I')} filed more than one bankruptcy. "
                     "The case must appear only once."))

            if opts.get("assert_unverifiable"):
                grounds.append(
                    ("unverifiable", "**Unverifiable source.** U.S. Bankruptcy Courts generally do not furnish data to, or verify data for, "
                     "consumer reporting agencies. If you cannot verify this item through a reliable source, "
                     "15 U.S.C. § 1681i(a)(5)(A) requires its prompt deletion."))

            for other in opts.get("other_grounds", {}).get(bureau, []):
                grounds.append(("other", f"**Other.** {other}"))

            if idt_ok:
                # D3: the identity theft block replaces the public-record and obsolete letters for this bureau.
                r["identity_theft"] = True
                r["reserve"] = r["obsolete"] + grounds
                r["obsolete"] = []
            elif r["obsolete"]:
                # D1: ask to delete, do not also ask to correct; keep the rest as reserve grounds for the follow-up.
                r["reserve"] = grounds
            else:
                r["public_record"] = grounds

        # Tradelines included in the bankruptcy
        if r["identity_theft"]:
            if rep.get("tradelines"):
                review.append(f"{B}: identity theft block requested for the bankruptcy record, so tradelines were not evaluated "
                              "as discharged debts. List fraudulent accounts in options.identity_theft_accounts.")
        else:
            matched = []
            for t in rep.get("tradelines", []):
                key_acct, key_name = norm(t.get("account_last4")), norm(t.get("creditor"))
                orig = norm(t.get("original_creditor"))
                match = next((s for s in scheduled if
                              (key_acct and norm(s.get("account_last4")) == key_acct) or
                              (key_name and (norm(s.get("name")) in key_name or key_name in norm(s.get("name")))) or
                              (orig and (norm(s.get("name")) in orig or orig in norm(s.get("name"))))), None)
                if not match:
                    continue
                label = f"{t.get('creditor')} ending {t.get('account_last4')}"
                # A collector or debt buyer may report an excluded debt under its own name and number, so also
                # test the original creditor and the matched scheduled creditor against the exclusions.
                names = [x for x in (key_name, orig, norm(match.get("name"))) if x]
                excl = next((x for (a, n), x in excluded.items()
                             if (a and a in (key_acct, norm(match.get("account_last4")))) or (n and any(n in s for s in names))), None)
                dofd = d(t.get("date_of_first_delinquency"))
                expires = add_years(dofd + timedelta(days=180), 7) if dofd else None
                if excl:
                    review.append(f"{B}: {label} is excluded/reaffirmed ({excl.get('reason', 'see schedule')}): not disputed automatically."
                                  + (f" Its date of first delinquency suggests the 7-year period ended {fmt(expires)}; some debts "
                                     "(for example student loans) have special rules, so confirm before disputing." if expires and today >= expires else ""))
                    continue
                opened = d(t.get("date_opened"))
                if opened and filed and opened > filed:
                    review.append(f"{B}: {label} opened after the petition date: likely post-petition debt, not disputed.")
                    continue
                if expires and today >= expires:
                    r["obsolete_tradelines"].append({**t, "dofd": dofd, "period_start": dofd + timedelta(days=180), "expires": expires})
                    continue
                if disposition == "discharged":
                    matched.append(t)

            # Discharged accounts still reported as owing
            for t in matched:
                problems = []
                if money(t.get("balance")) > 0:
                    problems.append(f"balance of ${money(t['balance']):,.2f}")
                if money(t.get("past_due")) > 0:
                    problems.append(f"past-due amount of ${money(t['past_due']):,.2f}")
                if not any(w in str(t.get("status", "")).lower() for w in BK_STATUS_WORDS):
                    problems.append(f"status \"{t.get('status')}\" instead of \"Discharged in Bankruptcy\"")
                if money(t.get("balance")) > 0:
                    for o in matched:
                        if o is t or money(o.get("balance")) <= 0:
                            continue
                        if t.get("original_creditor") and norm(o.get("creditor")) == norm(t.get("original_creditor")):
                            problems.append(f"a balance also reported by the original creditor, {o.get('creditor')}, so one discharged "
                                            "debt is counted twice")
                        elif o.get("original_creditor") and norm(o.get("original_creditor")) == norm(t.get("creditor")):
                            problems.append(f"a balance also reported by debt buyer or collector {o.get('creditor')}, so one discharged "
                                            "debt is counted twice")
                lates = t.get("late_payments_after_petition") or []
                if lates:
                    label = f"{t.get('creditor')} ending {t.get('account_last4')}"
                    if opts.get("assert_post_petition_lates"):
                        problems.append(f"late payments for {', '.join(lates)}, after the {fmt(filed)} petition date, on a debt "
                                        "included in the bankruptcy")
                        review.append(f"{B}: {label}: post-petition late-payment ground asserted because "
                                      "options.assert_post_petition_lates is true. This argument is contested in case law; confirm it fits.")
                        if chapter == "13":
                            review.append(f"{B}: {label}: Chapter 13 case. Post-petition late reporting is especially contested where "
                                          "the debtor kept paying under the plan or outside it. Attorney must confirm.")
                    else:
                        review.append(f"{B}: {label} reports late payments after the petition date ({', '.join(lates)}). Not asserted: "
                                      "this argument is contested in case law (especially in Chapter 13 or when the debtor kept paying). "
                                      "Set options.assert_post_petition_lates=true only after attorney review.")
                if problems:
                    r["tradelines"].append({**t, "problems": problems})
            if rep.get("tradelines") and disposition in ("open", "filed"):
                r["notes"].append("Case not yet discharged: tradelines not evaluated as discharged debts.")

        if recs and not (r["obsolete"] or r["public_record"] or r["identity_theft"]):
            keep_until = add_years(filed, 7 if chapter == "13" else 10) if filed else None
            r["notes"].append("Bankruptcy appears ACCURATELY reported: no good-faith ground to dispute it" +
                              (f"; expected to age off around {fmt(keep_until)}." if keep_until else "."))
        results[bureau] = r
    return results, review


def vendor_grounds(results, bureaus=None):
    """Public-record grounds to put to a public-records vendor (bureau-specific kinds left out).

    If any report names LexisNexis as its source, only those reports' grounds are used;
    otherwise (options.dispute_vendor) grounds from every bureau are used."""
    seen, out = set(), []
    sourced = any(r.get("vendor_source") for r in results.values())
    for b, r in results.items():
        if bureaus is not None:
            if b not in bureaus:
                continue
        elif sourced and not r.get("vendor_source"):
            continue  # a bureau named the vendor: use only that bureau's grounds
        for kind, text in r["obsolete"] + r["public_record"]:
            if kind in ("duplicate", "ch13_policy", "other"):
                continue
            key = (kind, text)
            if key not in seen:
                seen.add(key)
                out.append((kind, f"As reported by {name_of(b)}: {text}"))
    return out


# ---------------------------------------------------------------- rendering
def letterhead(case):
    if signer(case) == "attorney":
        a = case["attorney"]
        return [f"**{a['firm']}**", f"{a['address']} | {a['city_state_zip']}",
                f"{a['phone']} | {a['email']} | Bar No. {a.get('bar', '[____]')}", ""]
    c = case["client"]
    contact = " | ".join(x for x in (c.get("phone"), c.get("email")) if x)
    return [f"**{c['name']}**", c["address"]] + ([contact] if contact else []) + [""]


header = letterhead  # name used by generate_followup.py


def mail_line():
    return ["**VIA CERTIFIED MAIL, RETURN RECEIPT REQUESTED. No. [tracking #]**", ""]


def address_block(recipient, attn=None):
    n, l1, l2 = RECIPIENTS[recipient]
    return [n] + ([attn] if attn else []) + [l1, l2, ""]


def consumer_lines(case, recipient):
    c = case["client"]
    rep = case.get("credit_reports", {}).get(recipient) or {}
    return [f"**Consumer:** {c['name']}  ",
            f"**Date of Birth:** {c.get('dob', '[DOB]')}  **SSN (last 4):** XXX-XX-{c.get('ssn_last4', '####')}  ",
            f"**Current Address:** {c['address']}  ",
            f"**Report/File No.:** {rep.get('report_number', '[n/a]')}  "]


def case_line(case, suffix=""):
    court = case["bankruptcy_court_record"]
    return (f"**Bankruptcy:** Chapter {court.get('chapter')}, Case No. {court.get('case_number')}, "
            f"U.S. Bankruptcy Court, {court.get('district', '[District]')}{suffix}  ")


def opening(case):
    c = case["client"]
    return v(case,
             f"This firm represents {c['name']} (\"my client\"). My client's signed authorization, a copy of government-issued "
             "identification, and proof of address are enclosed as **Exhibit A**. Please send all correspondence about this dispute "
             "to my client and to this office.",
             "A copy of my government-issued identification and proof of my current address are enclosed as **Exhibit A**.")


def signature_block(case, enclosures, cc=True):
    c = case["client"]
    if signer(case) == "attorney":
        a = case["attorney"]
        L = ["Sincerely,", "", "", "______________________________", f"{a['name']}, Esq.  ", f"{a['firm']}  ",
             f"Counsel for {c['name']}", "", "**Enclosures:** " + "; ".join(enclosures)]
        return L + (["", f"cc: {c['name']}"] if cc else [])
    return ["Sincerely,", "", "", "______________________________", c["name"], "", "**Enclosures:** " + "; ".join(enclosures)]


def reinvestigation_request(case, recipient):
    n = reinvestigation_days(case, recipient)
    why = (v(case, " This dispute follows my client's free annual file disclosure.", " This dispute follows my free annual file disclosure.")
           if n == 45 else "")
    return f"Complete a reasonable reinvestigation within **{n} days** of receipt, as required by {deadline_basis(case, recipient)}.{why}"


def closing(case):
    return ["You are required to follow reasonable procedures to assure maximum possible accuracy (15 U.S.C. § 1681e(b)). "
            "Willful or negligent noncompliance gives rise to civil liability under 15 U.S.C. §§ 1681n and 1681o. "
            + v(case, "My client reserves", "I reserve") +
            f" all rights and remedies, including complaints to the Consumer Financial Protection Bureau and the {case.get('state', '[State]')} "
            "Attorney General.", "", "Thank you for your prompt attention to this matter.", ""]


def standard_requests(case, recipient, actions):
    who = v(case, "my client", "me")
    L = ["**Requested Action**", "", f"1. {reinvestigation_request(case, recipient)}"]
    L += [f"{i}. {x}" for i, x in enumerate(actions, 2)]
    n = len(actions) + 2
    L += [f"{n}. Send {who} written results and an updated copy of the consumer report, as required by 15 U.S.C. § 1681i(a)(6).",
          f"{n+1}. If any item is verified, provide a description of the procedure used, including the business name and address of "
          "any furnisher of information contacted and its telephone number, if reasonably available, as required by "
          "15 U.S.C. §§ 1681i(a)(6)(B)(iii) and 1681i(a)(7).",
          f"{n+2}. Do not reinsert any deleted item without complying with 15 U.S.C. § 1681i(a)(5)(B).", ""]
    return L


def disclosure_letter(case, recipient, letter_date):
    c = case["client"]
    L = letterhead(case) + [fmt(letter_date), ""] + mail_line() + address_block(recipient) + [
        "**Re: Request for Complete File Disclosure Under 15 U.S.C. § 1681g(a)**  "] + consumer_lines(case, recipient) + [
        "", "To Whom It May Concern:", "", opening(case), "",
        v(case, f"On behalf of my client, and under 15 U.S.C. § 1681g(a), I request", "Under 15 U.S.C. § 1681g(a), I request") +
        f" a complete disclosure of all information in {v(case, 'my client' + chr(39) + 's', 'my')} consumer file at the time of "
        "this request, including:", "",
        "1. All information in the file, including public record information and specifically any bankruptcy record (§ 1681g(a)(1));",
        "2. The sources of that information (§ 1681g(a)(2)), including the name of any public-records vendor or other entity that "
        "supplied any bankruptcy record;",
        "3. The identity of each person that procured a consumer report during the one-year period before this request, or the "
        "two-year period for employment purposes (§ 1681g(a)(3)).", "",
        v(case, f"Please send the disclosure to my client at {c['address']}.", "Please send the disclosure to me at the address above."), ""]
    L += signature_block(case, ["Exhibit A: " + v(case, "Client authorization, ", "") + "copy of government-issued photo ID, proof of current address"])
    return "\n".join(L)


def obsolete_letter(case, bureau, r, letter_date):
    court = case["bankruptcy_court_record"]
    L = letterhead(case) + [fmt(letter_date), ""] + mail_line() + address_block(bureau) + [
        "**Re: Dispute of Obsolete Information: Request for Deletion Under 15 U.S.C. §§ 1681c and 1681i**  "] + \
        consumer_lines(case, bureau) + [case_line(case), "", "To Whom It May Concern:", "", opening(case), "",
                                        v(case, "My client disputes", "I dispute") +
                                        " the following information because it is obsolete and can no longer lawfully be reported:", ""]
    if r["obsolete"]:
        L += ["**A. Bankruptcy Public Record**", ""] + [f"- {t}" for _, t in r["obsolete"]] + [""]
    if r["obsolete_tradelines"]:
        L += [f"**{'B' if r['obsolete'] else 'A'}. Accounts Past the Seven-Year Reporting Period**", "",
              "Under 15 U.S.C. § 1681c(a)(4) and (a)(5), an account placed for collection or charged to profit and loss, and any other "
              "adverse item of information, may not be reported once it antedates the report by more than seven years. Under 15 U.S.C. "
              "§ 1681c(c)(1), for a delinquent account placed for collection, charged to profit and loss, or subjected to any similar "
              "action, that seven-year period begins upon the expiration of the 180-day period beginning on the date the delinquency "
              "that immediately preceded that action began. Each account below is past that period:", ""]
        L += [f"- {t.get('creditor')}, account ending {t.get('account_last4')}: delinquency began {fmt(t['dofd'])}; the seven-year "
              f"period began {fmt(t['period_start'])} and ended {fmt(t['expires'])}." for t in r["obsolete_tradelines"]]
        L.append("")
    actions = []
    if r["obsolete"]:
        actions.append(f"**Delete** the bankruptcy public record from {v(case, 'my client' + chr(39) + 's', 'my')} file "
                       "(15 U.S.C. § 1681i(a)(5)(A)).")
    if r["obsolete_tradelines"]:
        actions.append("**Delete** each account listed above (15 U.S.C. § 1681i(a)(5)(A)).")
    L += standard_requests(case, bureau, actions)
    L += ["Willful or negligent noncompliance gives rise to civil liability under 15 U.S.C. §§ 1681n and 1681o. "
          + v(case, "My client reserves", "I reserve") + " all rights and remedies.", "",
          "Thank you for your prompt attention to this matter.", ""]
    enc = ["Exhibit A: " + v(case, "Client authorization, ", "") + "ID, proof of address",
           "Exhibit B: Court docket or notice of filing showing the filing date" if r["obsolete"] else None,
           "Exhibit D: Credit report with disputed items highlighted"]
    L += signature_block(case, [e for e in enc if e])
    return "\n".join(L)


def record_table(case, recs):
    court = case["bankruptcy_court_record"]
    L = ["| | Case Number | Chapter | Date Filed | Status | Status Date |", "|---|---|---|---|---|---|"]
    for i, pr in enumerate(recs, 1):
        L.append(f"| As reported{f' ({i})' if len(recs) > 1 else ''} | {pr.get('case_number', '')} | {pr.get('chapter', '')} | "
                 f"{fmt(d(pr.get('filing_date'))) if pr.get('filing_date') else ''} | {pr.get('status', '')} | "
                 f"{fmt(d(pr.get('status_date'))) if pr.get('status_date') else ''} |")
    L.append(f"| Court record | {court.get('case_number', '')} | {court.get('chapter', '')} | {fmt(d(court.get('filing_date')))} | "
             f"{court.get('disposition', '')} | {fmt(d(court.get('disposition_date')))} |")
    return L + [""]


def public_record_letter(case, bureau, r, letter_date):
    court = case["bankruptcy_court_record"]
    L = letterhead(case) + [fmt(letter_date), ""] + mail_line() + address_block(bureau) + [
        "**Re: Dispute of Inaccurate Bankruptcy Public Record Under 15 U.S.C. §§ 1681e(b) and 1681i**  "] + \
        consumer_lines(case, bureau) + [case_line(case), "", "To Whom It May Concern:", "", opening(case), "",
                                        "Under 15 U.S.C. § 1681i, " + v(case, "my client disputes", "I dispute") +
                                        " the bankruptcy public record in " + v(case, "my client's", "my") + " consumer file.", ""]
    L += record_table(case, r["records"])
    L += ["**Specific errors**", ""] + [f"- {t}" for _, t in r["public_record"]] + [""]
    kinds = {k for k, _ in r["public_record"]}
    file_of = v(case, "my client's", "my")
    actions = ["Review and consider the enclosed court records in your reinvestigation, as required by 15 U.S.C. § 1681i(a)(4), rather "
               "than simply re-confirming the existing data with its original source."]
    if kinds & set(DELETE_KINDS):
        actions.append(f"**Delete** the bankruptcy public record from {file_of} file (15 U.S.C. § 1681i(a)(5)(A)).")
    else:
        if "duplicate" in kinds:
            actions.append("**Delete** every duplicate listing so the case appears only once.")
        if kinds - {"duplicate"}:
            actions.append("**Correct** the bankruptcy public record to match the enclosed court records, or delete it if it cannot be "
                           "verified (15 U.S.C. § 1681i(a)(5)(A)).")
    L += standard_requests(case, bureau, actions)
    L += closing(case)
    disp = str(court.get("disposition", "")).lower()
    enc = ["Exhibit A: " + v(case, "Client authorization, ", "") + "ID, proof of address", "Exhibit B: Court docket/records",
           ("Exhibit C: Discharge order" if disp.startswith("discharg") else
            "Exhibit C: Dismissal or withdrawal order" if disp.startswith(("dismiss", "withdraw")) else None),
           "Exhibit D: Credit report with disputed items highlighted"]
    L += signature_block(case, [e for e in enc if e])
    return "\n".join(L)


def discharged_accounts_letter(case, bureau, r, letter_date):
    court = case["bankruptcy_court_record"]
    who = v(case, "my client", "I")
    L = letterhead(case) + [fmt(letter_date), ""] + mail_line() + address_block(bureau) + [
        "**Re: Dispute of Inaccurate Reporting on Accounts Discharged in Bankruptcy**  "] + \
        consumer_lines(case, bureau) + [case_line(case, f"; discharged {fmt(d(court.get('disposition_date')))}"), "",
                                        "To Whom It May Concern:", "", opening(case), "",
                                        f"The debts below were scheduled in the bankruptcy and discharged by order entered "
                                        f"{fmt(d(court.get('disposition_date')))} (Exhibit C). Under 11 U.S.C. § 524(a), a discharge voids any "
                                        "judgment to the extent it determines personal liability on a discharged debt and operates as an "
                                        "injunction against any act to collect the debt as a personal liability of the debtor. Each account "
                                        f"below is nevertheless reported in a way that indicates {who} still "
                                        + v(case, "owes", "owe") + " the debt.", ""]
    for t in r["tradelines"]:
        L += [f"**{t.get('creditor')}, account ending {t.get('account_last4')}**", "",
              "- Reported: " + "; ".join(t["problems"]) + ".",
              "- Should report: $0 balance, $0 past due, status \"Discharged in Bankruptcy\".", ""]
    L += ["Reporting a balance or past-due amount on a discharged debt is inaccurate and fails your duty to follow reasonable "
          "procedures to assure maximum possible accuracy (15 U.S.C. § 1681e(b)).", ""]
    if bureau in NATIONWIDE and str(court.get("chapter", "")).strip() == "7":
        L += ["Your company has also been on notice of this specific issue since the court-approved settlement order in White v. "
              "Experian Information Solutions, Inc., No. 8:05-cv-01070-DOC-MLG (C.D. Cal. Aug. 19, 2008), under which Equifax, "
              "Experian, and TransUnion follow procedures that presume certain categories of pre-bankruptcy debts (such as accounts "
              "in collection, charged off, or seriously past due) were discharged after a Chapter 7 discharge, and report them as "
              "discharged with no balance owed.", ""]
    actions = ["**Update** each account to a $0 balance, $0 past due, and a status of \"Discharged in Bankruptcy\", or delete any "
               "account whose accuracy cannot be verified (15 U.S.C. § 1681i(a)(5)(A)).",
               "Forward all relevant information about this dispute, including the enclosures, to each furnisher, as required by "
               "15 U.S.C. § 1681i(a)(2)."]
    L += standard_requests(case, bureau, actions)
    L += closing(case)
    L += signature_block(case, ["Exhibit A: " + v(case, "Client authorization, ", "") + "ID, proof of address",
                                "Exhibit B: Schedules listing each creditor (SSN redacted)", "Exhibit C: Discharge order",
                                "Exhibit D: Credit report with each account highlighted"])
    return "\n".join(L)


def identity_theft_letter(case, bureau, r, letter_date):
    c, opts = case["client"], case.get("options", {})
    my = v(case, "my client's", "my")
    L = letterhead(case) + [fmt(letter_date), ""] + mail_line() + address_block(bureau, "Attn: Fraud Department / Identity Theft Block") + [
        "**Re: Request to Block Information Resulting From Identity Theft Under 15 U.S.C. § 1681c-2**  ",
        f"**FTC Identity Theft Report No.:** {opts.get('ftc_report_number')}  "] + consumer_lines(case, bureau) + [
        "", "To Whom It May Concern:", "", opening(case), "",
        v(case, "My client is a victim of identity theft. The following bankruptcy record in my client's file resulted from identity "
                "theft. My client did not file it, authorize it, or participate in it:",
          "I am a victim of identity theft. The following bankruptcy record in my file resulted from identity theft. I did not file "
          "it, authorize it, or participate in it:"), ""]
    for pr in r["records"]:
        L.append(f"- Case No. {pr.get('case_number', '[case no.]')}, Chapter {pr.get('chapter', '[ ]')}, "
                 f"filed {fmt(d(pr.get('filing_date')))}, court: {pr.get('court', '[court]')}, reported status: {pr.get('status', '[ ]')}")
    accts = opts.get("identity_theft_accounts") or []
    if accts:
        L += ["", "The following accounts also resulted from identity theft:", ""]
        L += [f"- {a.get('creditor')}, account ending {a.get('account_last4')}" for a in accts]
    L += ["", f"Under 15 U.S.C. § 1681c-2(a), please block this information from {my} file not later than four business days after "
          "you receive this request. Enclosed are:", "",
          f"1. Appropriate proof of {my} identity (Exhibit A);",
          "2. A copy of the FTC Identity Theft Report (Exhibit B);",
          "3. This letter, which identifies the information that resulted from identity theft; and",
          f"4. The signed statement below that the information does not relate to any transaction by {v(case, 'my client', 'me')}.", "",
          "As required by 15 U.S.C. § 1681c-2(b), please promptly notify the furnisher or source of this information that it may be "
          "the result of identity theft, that an identity theft report has been filed, that a block has been requested, and of the "
          "effective dates of the block."]
    if bureau in NATIONWIDE:
        L += ["", f"Please also place an initial fraud alert in {my} file under 15 U.S.C. § 1681c-1(a)."]
    L += ["", "**Consumer statement**", "",
          f"I, {c['name']}, state that the information identified above is not information relating to any transaction by me.", "",
          "______________________________  Date: ____________  ", c["name"], ""]
    L += signature_block(case, ["Exhibit A: " + v(case, "Client authorization, ", "") + "ID, proof of address",
                                "Exhibit B: FTC Identity Theft Report (and police report, if any)",
                                "Exhibit C: Evidence the case is not " + v(case, "my client's", "mine") + " (if any)",
                                "Exhibit D: Credit report with the record highlighted"])
    return "\n".join(L)


def vendor_letter(case, grounds, letter_date, source_note=None):
    court = case["bankruptcy_court_record"]
    who = v(case, "my client", "me")
    intro = (source_note or v(case, "My client understands that your company may supply bankruptcy public-record data to consumer "
                                    "reporting agencies.",
                              "I understand that your company may supply bankruptcy public-record data to consumer reporting agencies."))
    L = letterhead(case) + [fmt(letter_date), ""] + mail_line() + address_block("lexisnexis") + [
        "**Re: Dispute of Bankruptcy Record Under 15 U.S.C. § 1681i**  "] + consumer_lines(case, "lexisnexis") + [
        case_line(case), "", "To Whom It May Concern:", "", opening(case), "", intro, "",
        v(case, "On behalf of my client, I request", "I request") + " a copy of " + v(case, "my client's", "my") +
        " consumer file under 15 U.S.C. § 1681g(a), and " + v(case, "my client disputes", "I dispute") + " the bankruptcy record "
        f"for Case No. {court.get('case_number')} for the following reasons:", ""]
    L += [f"- {t}" for _, t in grounds] + [""]
    deleting = any(k in DELETE_KINDS for k, _ in grounds)
    L += ["**Requested Action**", "",
          f"1. {reinvestigation_request(case, 'lexisnexis')} Reinvestigate against the actual court docket and review and consider "
          "the enclosed documents (15 U.S.C. § 1681i(a)(4)).",
          ("2. **Delete** the record" if deleting else "2. **Correct** the record to match the court docket, or delete it if it cannot be verified")
          + " (15 U.S.C. § 1681i(a)(5)(A)).",
          "3. Notify each consumer reporting agency to which you supplied this record of the correction or deletion.",
          f"4. Send {who} written results and a copy of the consumer file (15 U.S.C. § 1681i(a)(6)).", "",
          "Willful or negligent noncompliance gives rise to civil liability under 15 U.S.C. §§ 1681n and 1681o.", "",
          "Thank you for your prompt attention to this matter.", ""]
    L += signature_block(case, ["Exhibit A: " + v(case, "Client authorization, ", "") + "ID, proof of address",
                                "Exhibit B: Court docket/orders", "Exhibit D: Credit bureau report showing the record"])
    return "\n".join(L)


def furnisher_letter(case, creditor, items, letter_date):
    c, court = case["client"], case["bankruptcy_court_record"]
    addr = items[0].get("furnisher_address") or "[Furnisher dispute address from credit report]"
    bureaus = ", ".join(sorted({name_of(i["_bureau"]) for i in items}))
    probs = list(dict.fromkeys(p for i in items for p in i["problems"]))
    L = letterhead(case) + [fmt(letter_date), ""] + mail_line() + [creditor, *addr.split("\n"), "",
                                                                   f"**Re: Direct Dispute Under 15 U.S.C. § 1681s-2(a)(8) and 12 C.F.R. § 1022.43: "
                                                                   f"Account ending {items[0].get('account_last4')}**  ",
                                                                   f"**Consumer:** {c['name']}, SSN last 4: {c.get('ssn_last4')}, {c['address']}  ",
                                                                   f"**Bankruptcy:** Chapter {court.get('chapter')}, Case No. {court.get('case_number')}, {court.get('district')}; "
                                                                   f"discharged {fmt(d(court.get('disposition_date')))}", "",
                                                                   "To Whom It May Concern:", "",
                                                                   v(case, f"This firm represents {c['name']}; my client's signed authorization is enclosed. ", "") +
                                                                   "This account was scheduled in the above bankruptcy and the debt was discharged "
                                                                   "(Discharge Order enclosed). You report it to " + bureaus + " with: " + "; ".join(probs) + ".", "",
                                                                   "This reporting is inaccurate: the debt was discharged and no balance is owed. The enclosed discharge order, "
                                                                   "schedule, and credit report excerpt substantiate this dispute.", "",
                                                                   "Under 12 C.F.R. § 1022.43(e), you must conduct a reasonable investigation, review all relevant information "
                                                                   "provided with this notice, complete the investigation and report the results to "
                                                                   + v(case, "my client", "me") + " within the period that would apply to a consumer reporting agency under "
                                                                   "15 U.S.C. § 1681i(a)(1), and, if the information is inaccurate, promptly notify each consumer reporting "
                                                                   "agency to which you furnished it and provide the correction.", "",
                                                                   "Please report this account to every agency with a $0 balance, $0 past due, and a status of \"Discharged in "
                                                                   "Bankruptcy\", and confirm the correction in writing. Any attempt to collect this discharged debt as a "
                                                                   "personal liability violates the discharge injunction, 11 U.S.C. § 524(a)(2).", ""]
    L += signature_block(case, [v(case, "Client authorization", "Copy of ID and proof of address"), "Discharge Order",
                                "Schedule listing this account (SSN redacted)", "Credit report excerpt"], cc=False)
    return "\n".join(L)


def authorization(case, recipients):
    c, a = case["client"], case["attorney"]
    names = [name_of(x) for x in recipients]
    listed = ", ".join(names[:-1]) + (", and " if len(names) > 2 else " and ") + names[-1] if len(names) > 1 else (names[0] if names else "consumer reporting agencies")
    return "\n".join(["# Client Authorization to Dispute Credit Information", "",
                      f"I, **{c['name']}**, DOB {c.get('dob', '[DOB]')}, SSN last 4 {c.get('ssn_last4', '[####]')}, residing at {c['address']}, "
                      f"authorize **{a['name']}, Esq.** and **{a['firm']}** to act on my behalf in disputing information in my consumer files "
                      f"with {listed}, and with any furnisher of information, and to receive information about those disputes. "
                      "I confirm that, to the best of my knowledge, the statements in the enclosed disputes are true and correct.", "", "",
                      "______________________________  Date: ____________  ", c["name"]])


def to_docx(md_text, path):
    try:
        import docx
        from docx.shared import Pt
    except ImportError:
        return False
    doc = docx.Document()
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(12)
    for line in md_text.split("\n"):
        line = line.rstrip()
        if line.startswith("# "):
            doc.add_heading(line[2:], level=1)
            continue
        if re.fullmatch(r"\|[-| ]+\|", line):
            continue  # markdown table separator row
        style = None
        if line.startswith("- "):
            line, style = line[2:], "List Bullet"
        elif line.startswith("|"):
            line = "   ".join(x.strip() for x in line.strip("|").split("|"))
        p = doc.add_paragraph(style=style)
        p.paragraph_format.space_after = Pt(0 if not line else 4)
        for i, chunk in enumerate(re.split(r"\*\*", line)):
            if chunk:
                p.add_run(chunk).bold = i % 2 == 1
    doc.save(str(path))
    return True


def write(md, path):
    path.write_text(md)
    return to_docx(md, path.with_suffix(".docx"))


# ---------------------------------------------------------------- main
def disclosure_targets(case):
    req = case.get("options", {}).get("request_disclosure")
    if req is True:
        return list(RECIPIENTS)
    targets = [x.lower() for x in req] if isinstance(req, list) else []
    have = {b.lower() for b in case.get("credit_reports", {})}
    targets += [b for b in NATIONWIDE if b not in have]  # D5: missing nationwide reports
    return [b for b in RECIPIENTS if b in targets]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("case")
    ap.add_argument("--out", required=True)
    ap.add_argument("--today")
    args = ap.parse_args()
    case = json.loads(Path(args.case).read_text())
    if signer(case) not in ("attorney", "consumer"):
        sys.exit("case.json signer must be \"attorney\" or \"consumer\".")
    for k in (["attorney"] if signer(case) == "attorney" else []) + ["client", "bankruptcy_court_record", "credit_reports"]:
        if k not in case:
            sys.exit(f"case.json missing required section: {k}")
    today = d(args.today) or date.today()
    mail = d(case.get("mail_date")) or today
    out = Path(args.out)
    (out / "letters").mkdir(parents=True, exist_ok=True)
    opts = case.get("options", {})

    results, review = analyze(case, today)
    docx_ok = True
    files = []          # (relative path, recipient, kind)
    furnishers = {}

    def emit(md, name, recipient, kind):
        nonlocal docx_ok
        docx_ok &= write(md, out / "letters" / f"{name}.md")
        files.append((f"letters/{name}.md", recipient, kind))

    for b in disclosure_targets(case):
        emit(disclosure_letter(case, b, mail), f"{b}-file-disclosure", b, "disclosure")
    for bureau, r in results.items():
        if r["obsolete"] or r["obsolete_tradelines"]:
            emit(obsolete_letter(case, bureau, r, mail), f"{bureau}-obsolete", bureau, "dispute")
        if r["public_record"]:
            emit(public_record_letter(case, bureau, r, mail), f"{bureau}-public-record", bureau, "dispute")
        if r["tradelines"]:
            emit(discharged_accounts_letter(case, bureau, r, mail), f"{bureau}-discharged-accounts", bureau, "dispute")
        if r["identity_theft"]:
            emit(identity_theft_letter(case, bureau, r, mail), f"{bureau}-identity-theft-block", bureau, "identity_theft")
        for t in r["tradelines"]:
            furnishers.setdefault((t.get("creditor"), t.get("account_last4")), []).append({**t, "_bureau": bureau})

    vendor_wanted = opts.get("dispute_vendor") or any(r["vendor_source"] for r in results.values())
    vg = vendor_grounds(results) if vendor_wanted else []
    vendor_src = next(((b, r["vendor_source"]) for b, r in results.items() if r["vendor_source"]), None)
    if vg:
        note = (v(case, f"{name_of(vendor_src[0])} identified your company as the source of the bankruptcy record in my client's file.",
                  f"{name_of(vendor_src[0])} identified your company as the source of the bankruptcy record in my file.")
                if vendor_src else None)
        emit(vendor_letter(case, vg, mail, note), "lexisnexis-vendor-dispute", "lexisnexis", "dispute")
    for (cred, acct), items in furnishers.items():
        emit(furnisher_letter(case, cred, items, mail), f"furnisher-{slug(cred)}-{acct}", None, "furnisher")
    recipients_used = list(dict.fromkeys(rc for _, rc, _ in files if rc))
    if files and signer(case) == "attorney":
        docx_ok &= write(authorization(case, recipients_used), out / "client-authorization.md")

    # analysis report
    court = case["bankruptcy_court_record"]
    A = [f"# Dispute Analysis: {case['client']['name']}", "", f"_Generated {fmt(today)}. " + v(case, "Attorney review required before mailing._", "Review every statement before mailing._"), "",
         f"**Signer:** {signer(case)}  ",
         f"**Court record:** Chapter {court.get('chapter')}, Case {court.get('case_number')}, {court.get('district')}; "
         f"filed {fmt(d(court.get('filing_date')))}; {court.get('disposition')} {fmt(d(court.get('disposition_date')))}", ""]
    for bureau, r in results.items():
        A.append(f"## {name_of(bureau)}")
        A.append(f"- Report source: {(case['credit_reports'].get(bureau) or {}).get('report_source', 'other')}; "
                 f"reinvestigation period {reinvestigation_days(case, bureau)} days.")
        if len(r["records"]) > 1:
            A.append(f"- {len(r['records'])} public records reported for this bureau.")
        A += [f"- Obsolete ground ({k}): {t}" for k, t in r["obsolete"]]
        A += [f"- Obsolete account: {t.get('creditor')} ending {t.get('account_last4')}: 7-year period ended {fmt(t['expires'])}"
              for t in r["obsolete_tradelines"]]
        A += [f"- Public record ground ({k}): {t}" for k, t in r["public_record"]]
        if r["identity_theft"]:
            A.append("- Identity theft block requested (options.identity_theft_confirmed with FTC report number).")
        A += [f"- Reserve ground, not in a letter ({k}): {t}" for k, t in r["reserve"]]
        A += [f"- Tradeline: {t.get('creditor')} ending {t.get('account_last4')}: " + "; ".join(t["problems"]) for t in r["tradelines"]]
        if r["vendor_source"]:
            A.append(f"- Public record source: {r['vendor_source']}.")
        A += [f"- Note: {n}" for n in r["notes"]]
        A.append("")
    if any(r["obsolete"] or r["obsolete_tradelines"] or r["reserve"] for r in results.values()):
        A += ["**Reporting-period exceptions:** the time limits in 15 U.S.C. § 1681c(a)(1)-(5) do not apply to a report used for a credit "
              "transaction of $150,000 or more, life insurance with a face amount of $150,000 or more, or employment at an annual salary "
              "of $75,000 or more (15 U.S.C. § 1681c(b)), so an item past these periods may still lawfully appear in a report for "
              "those uses.", ""]
    if vendor_wanted and not vg:
        A += ["**Vendor dispute:** requested or a vendor source was found, but no public-record ground applies, so no vendor letter was written.", ""]
    A += ["## Needs attorney review (not included in any letter)", ""] + ([f"- {x}" for x in review] or ["- None"])
    A += ["", "## Generated files", ""] + [f"- {p}" for p, _, _ in files] + \
         (["- client-authorization.md"] if files and signer(case) == "attorney" else []) + \
         ([] if files else ["- (no letters: nothing to dispute in good faith)"])
    (out / "analysis.md").write_text("\n".join(A))

    # checklist
    received = mail + timedelta(days=MAIL_DAYS)
    C = ["# Mailing Checklist & Deadlines", "", f"Planned mail date: **{fmt(mail)}**", ""]
    C += ["- [ ] Attorney reviewed analysis.md and every letter" if signer(case) == "attorney" else
          "- [ ] Read analysis.md and every letter; confirm every statement is true",
          ("- [ ] Client signed client-authorization; attach ID + proof of address (Exhibit A)" if signer(case) == "attorney" else
           "- [ ] Attach copies of ID + proof of address (Exhibit A); send the letters yourself, not through a credit repair service"),
          "- [ ] Attach PACER docket / court records (Exhibit B) and discharge or dismissal order (Exhibit C); redact SSNs to the last 4",
          "- [ ] Highlight disputed items on each credit report (Exhibit D)",
          "- [ ] One issue per letter: mail each letter separately",
          "- [ ] Fill in furnisher dispute addresses (use the address on the credit report or the furnisher's stated dispute address, "
          "12 C.F.R. § 1022.43(c))",
          "- [ ] Send certified mail, return receipt; record tracking numbers"]
    if any(r["identity_theft"] for r in results.values()):
        C += ["- [ ] Identity theft: file at IdentityTheft.gov and attach the FTC Identity Theft Report; client signs the consumer "
              "statement in each block letter; also notify the U.S. Trustee's office and the bankruptcy court clerk"]
    C += ["", "**Addresses (verify every address before mailing)**", ""]
    for rc in recipients_used:
        n, l1, l2 = RECIPIENTS[rc]
        flag = (" (not confirmed: the innovis.com contact page lists Innovis Consumer Assistance, PO Box 530088, Atlanta, GA "
                "30353-0088 and the dispute page gives only online and phone options; confirm the dispute address before mailing)"
                if rc == "innovis" else "")
        C.append(f"- {n}, {l1}, {l2}{flag}")
    if not recipients_used:
        C.append("- (no agency letters)")
    C += ["", "| Deadline | Date |", "|---|---|"]
    for rc in recipients_used:
        kinds = {k for _, x, k in files if x == rc}
        if "dispute" in kinds:
            n = reinvestigation_days(case, rc)
            due = received + timedelta(days=n)
            C.append(f"| {name_of(rc)}: reinvestigation due ({n} days after receipt, assume {MAIL_DAYS} days mail; "
                     f"{deadline_basis(case, rc)}) | {fmt(due)} |")
            if n == 30:
                C.append(f"| {name_of(rc)}: extended deadline if the client sends relevant new information during the 30 days "
                         f"(45 days, 15 U.S.C. § 1681i(a)(1)(B)) | {fmt(received + timedelta(days=45))} |")
            C.append(f"| {name_of(rc)}: results notice due (5 business days after completion, 15 U.S.C. § 1681i(a)(6)(A)) | "
                     f"approx. {fmt(add_business_days(due, 5))} |")
        if "identity_theft" in kinds:
            C.append(f"| {name_of(rc)}: identity theft block due (4 business days after receipt, 15 U.S.C. § 1681c-2(a)) | "
                     f"{fmt(add_business_days(received, 4))} |")
        if kinds == {"disclosure"}:
            C.append(f"| {name_of(rc)}: file disclosure requested (no statutory dispute deadline) | {fmt(received)} received (est.) |")
    if furnishers:
        C.append(f"| Furnishers: direct dispute results due (same period as a bureau reinvestigation, 12 C.F.R. § 1022.43(e)(3); "
                 f"30 days assumed) | {fmt(received + timedelta(days=30))} |")
    C.append(f"| Pull fresh reports to confirm corrections | {fmt(mail + timedelta(days=60))} |")
    C += ["", "Business-day dates skip weekends only; federal holidays are not counted, so allow extra time.", ""]
    if any((case["credit_reports"].get(b) or {}).get("report_source") == "annual_free" for b in results):
        C += ["The 45-day period applies to a dispute made after receiving a free annual disclosure under 15 U.S.C. § 1681j(a) "
              "(15 U.S.C. § 1681j(a)(3)). Confirm the report actually was a § 1681j(a) annual disclosure; if unsure, calendar the "
              "30-day date.", ""]
    if furnishers:
        C += ["Furnisher letters are direct disputes under 15 U.S.C. § 1681s-2(a)(8) and 12 C.F.R. § 1022.43. A furnisher's duties under "
              "15 U.S.C. § 1681s-2(b) arise only after a bureau notifies it of the dispute under 15 U.S.C. § 1681i(a)(2), which is why "
              "the bureau letter is sent too: 15 U.S.C. § 1681s-2(c) bars private claims under §§ 1681n and 1681o for violations of "
              "§ 1681s-2(a), which includes direct disputes, but not for violations of § 1681s-2(b). "
              "A furnisher may decline a direct dispute it reasonably believes was prepared by a "
              "credit repair organization (12 C.F.R. § 1022.43(b)(2)) or that is frivolous or irrelevant (12 C.F.R. § 1022.43(f)).", ""]
    C += ["A bureau may end a reinvestigation it reasonably determines is frivolous or irrelevant, including for lack of information "
          "(15 U.S.C. § 1681i(a)(3)): include the identifying details and documents each letter lists.", ""]
    if any(r["obsolete"] or r["obsolete_tradelines"] for r in results.values()):
        C += ["Reporting limits do not apply to reports for a credit transaction with a principal amount of $150,000 or more, life "
              "insurance with a face amount of $150,000 or more, or employment at an annual salary of $75,000 or more "
              "(15 U.S.C. § 1681c(b)).", ""]
    C += ["If an item is verified without a description of the procedure, or no response arrives: run the follow-up, file a CFPB "
          "complaint, and evaluate FCRA claims (§§ 1681n and 1681o; suit is due by the earlier of 2 years after discovery of the "
          "violation or 5 years after it occurred, § 1681p)."]
    (out / "checklist.md").write_text("\n".join(C))

    print((out / "analysis.md").read_text())
    if not docx_ok:
        print("\n(note: python-docx not installed: only .md files written)")


if __name__ == "__main__":
    main()
